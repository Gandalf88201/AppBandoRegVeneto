#!/usr/bin/env python3
"""
aggiorna_bandi.py - aggiorna il dataset dei bandi della Regione Veneto e rigenera l'app HTML standalone.

Fonti (ognuna è una funzione indipendente: se una fallisce le altre proseguono):
  1. Portale Bandi Avvisi Concorsi (bandi.regione.veneto.it)  -> bandi (tutti) + avvisi pubblicati
  2. Calendario inviti FESR 2021-27 (venetocoesione.regione.veneto.it) -> bandi PREVISTI
  3. Calendari inviti FSE+ (stesso sito, CSV)                  -> bandi PREVISTI
  4. Foglio "Opportunità" di Veneto Coesione (Google Sheet CSV) -> atti FESR/INTERREG non presenti nel portale

Uso:
  python aggiorna_bandi.py                 # aggiornamento giornaliero (default)
  python aggiorna_bandi.py --max-details 400 --time-budget 1500   # passata settimanale più profonda
  python aggiorna_bandi.py --no-refresh-open --max-details 0      # solo liste e calendari (veloce)

Il dettaglio dei bandi chiusi più vecchi si completa gradualmente: ogni esecuzione ne scarica
al massimo --max-details (nuovi prima, poi i più recenti). Il risultato è in cache/detail_cache.json.
Uscite: data/bandi_veneto.json, data/bandi_veneto.xlsx (Excel), bandi_veneto.html (app standalone).
"""
import argparse, csv, html as htmllib, io, json, os, re, sys, time, datetime as dt
import threading
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup

PORTAL = "https://bandi.regione.veneto.it"
COESIONE = "https://venetocoesione.regione.veneto.it"
SHEET = ("https://docs.google.com/spreadsheets/d/1n4iPaRDMrc1p5jxpMvVSsTDn49S-P6OHKWyK1qp9VkU/"
         "gviz/tq?tqx=out:csv&sheet=WebService&range=A:W")
UA = "Mozilla/5.0 (compatible; bandi-veneto-dashboard/1.0)"
HERE = os.path.dirname(os.path.abspath(__file__))
TODAY = dt.date.today()

S = requests.Session()
S.headers.update({"User-Agent": UA})


def log(*a):
    print(time.strftime("%H:%M:%S"), *a, flush=True)


# ------------------------------------------------------------------ utilità
def get(url, **kw):
    kw.setdefault("timeout", 30)
    for k in range(3):
        try:
            r = S.get(url, **kw)
            if r.status_code == 200:
                return r
        except requests.RequestException:
            pass
        time.sleep(2 * (k + 1))
    raise RuntimeError(f"GET fallita: {url}")


def iso(s):
    m = re.match(r"\s*(\d\d)/(\d\d)/(\d{4})", s or "")
    return f"{m.group(3)}-{m.group(2)}-{m.group(1)}" if m else ""


def txt_date(s):
    m = re.match(r"\s*(\d\d/\d\d/\d{4})", s or "")
    return m.group(1) if m else ""


def importo(s):
    s = re.sub(r"[^\d,.]", "", s or "")
    try:
        v = float(s.replace(".", "").replace(",", "."))
    except ValueError:
        return None
    return v if v > 0 else None


def importo_txt(v):
    if not v:
        return "n.d."
    if v >= 1e6:
        return f"{v / 1e6:.1f}".replace(".", ",").rstrip("0").rstrip(",") + " M€"
    if v >= 1e3:
        return f"{round(v / 1e3)} k€"
    return f"{int(v)} €"


def fix(s):
    s = re.sub(r"\s+", " ", s or "").strip()
    return s.replace("Attivita'", "Attività").replace("Sanita'", "Sanità").replace("Mobilita'", "Mobilità")


def toks(s):
    return set(re.findall(r"[a-zà-ú0-9]{4,}", (s or "").lower()))


def jaccard(a, b):
    return len(a & b) / max(1, len(a | b))


MESI = {"gennaio": 1, "gen": 1, "febbraio": 2, "feb": 2, "marzo": 3, "mar": 3, "aprile": 4, "apr": 4,
        "maggio": 5, "mag": 5, "giugno": 6, "giu": 6, "luglio": 7, "lug": 7, "agosto": 8, "ago": 8,
        "settembre": 9, "set": 9, "sett": 9, "ottobre": 10, "ott": 10, "novembre": 11, "nov": 11,
        "dicembre": 12, "dic": 12}
NOMI_MESI = ["", "gen", "feb", "mar", "apr", "mag", "giu", "lug", "ago", "set", "ott", "nov", "dic"]


def parse_finestra(s):
    """'Ottobre – Novembre 2026', 'nov-26', 'Dicembre 2026' -> (inizio_iso, fine_iso) oppure (None, None)."""
    t = re.findall(r"[a-zà-ù]+|\d{2,4}", (s or "").lower())
    mesi = [MESI[x] for x in t if x in MESI]
    anni = [int(x) + (2000 if int(x) < 100 else 0) for x in t if x.isdigit() and len(x) in (2, 4)]
    if not mesi or not anni:
        return None, None
    a0, a1 = anni[0], anni[-1]
    start = dt.date(a0, mesi[0], 1)
    ey, em = a1, mesi[-1]
    end = (dt.date(ey + (em == 12), em % 12 + 1, 1) - dt.timedelta(days=1))
    return start.isoformat(), end.isoformat()


# ------------------------------------------------------------------ fonte 1: portale
def portale_lista(tipo, stato=""):
    rows, start, page = [], 0, 250
    while True:
        p = dict(cig="", parolaChiave="", tipoAttoTmp=tipo, destinatariTmp="", struttureTmp="0", categorieTmp="",
                 statoTmp=stato, materieTmp="", paginaIniziale="Elenco", sEcho="1", iColumns="1", sColumns="",
                 iDisplayStart=start, iDisplayLength=page, mDataProp_0="0")
        j = get(PORTAL + "/Public/GetListaAttiJson", params=p, headers={"X-Requested-With": "XMLHttpRequest"}).json()
        for (h,) in j["aaData"]:
            d = BeautifulSoup(h, "html.parser")
            idatto = int(d.find("span").text.strip())
            st, sc = d.select_one("span[class^=stato]"), d.select_one(".DtScadenza")
            a, stru = d.select_one("a[href*=Dettaglio]"), d.select_one(".selenium-field-struttura")
            rows.append(dict(id=idatto, tipo=int(tipo), stato=st.text.strip() if st else "",
                             scadenza=sc.text.strip() if sc else "", titolo=a.text.strip() if a else "",
                             struttura=stru.text.strip() if stru else ""))
        start += page
        if start >= j["recordsTotal"] or not j["aaData"]:
            break
        time.sleep(0.4)
    return rows


def portale_dettaglio(idatto):
    r = S.get(PORTAL + "/Public/Dettaglio", params={"idAtto": idatto}, timeout=25)
    r.raise_for_status()
    d = BeautifulSoup(r.text, "html.parser")
    for t in d(["script", "style"]):
        t.decompose()
    full = d.get_text("\n", strip=True)

    def after(label, stop):
        m = re.search(re.escape(label) + r"\n(.*?)\n(?:" + stop + ")", full, re.S)
        return re.sub(r"\s+", " ", m.group(1)).strip() if m else ""

    out = {"id": idatto}
    m = re.search(r"Pubblicazione:\n(\d\d/\d\d/\d{4})", full)
    out["pubblicazione"] = m.group(1) if m else ""
    out["oggetto"] = after("Oggetto", "Dati BUR")
    out["categoria"] = after("Categoria", "Importo|Struttura")
    out["importo"] = after("Importo", "Struttura")
    out["destinatari"] = after("Destinatari", "Materie")
    m = re.search(r"Materie\n(.*?)\n(?:Acquisti Verdi|NOTE:|CIG)", full, re.S)
    out["materie"] = [x.strip() for x in m.group(1).split("\n") if x.strip()] if m else []
    out["nattach"] = len(re.findall(r"Numero allegato:", full))
    out["_f"] = TODAY.isoformat()
    return out


FIN = {"Contributo", "Finanziamento", "Fondo di rotazione", "Riconoscimento", "Avviso di contributo", "- Non Selezionata -"}
GARE = {"Bando di gara", "Bando di gara comunitario", "Affidamento diretto", "Procedura negoziata"}
GARE_RE = re.compile(r"affidamento|procedura negoziata|\bgara\b|appalto|fornitura|incarico|concessione|noleggio|"
                     r"evidenza pubblica per il conferimento|servizio di|servizi di|manifestazione di interesse", re.I)


def classifica(tipo, cat, titolo, ha_dettaglio):
    if ha_dettaglio:
        if cat in GARE:
            return "Gare e affidamenti"
        if cat in FIN and (tipo == "Bando" or cat != "- Non Selezionata -"):
            return "Finanziamenti"
        return "Altro"
    return "Gare e affidamenti" if GARE_RE.search(titolo) else "Finanziamenti"


def riga_vuota():
    return {"id": "", "tipo": "", "stato_portale": "", "titolo": "", "struttura": "", "categoria": "", "gruppo": "",
            "scadenza": "", "scad_iso": "", "pubblicazione": "", "pub_iso": "", "importo": None, "importo_txt": "",
            "destinatari": "", "materie": "", "oggetto": "", "allegati": 0, "dettaglio": False, "link": "",
            "fonte": "", "apertura_prevista": ""}


def fonte_portale(args, cache):
    log("portale: liste")
    bandi = portale_lista("1")
    avvisi = portale_lista("2", stato="4")
    items = bandi + avvisi
    log(f"portale: {len(bandi)} bandi, {len(avvisi)} avvisi pubblicati")

    # quali dettagli scaricare: aperti (sempre, salvo --no-refresh-open), poi non in cache, più recenti prima
    aperti = [i for i in items if i["stato"] == "PUBBLICATO"]
    mancanti = [i for i in items if i["stato"] != "PUBBLICATO" and str(i["id"]) not in cache]
    todo = []
    if not args.no_refresh_open:
        todo += aperti
    else:
        todo += [i for i in aperti if str(i["id"]) not in cache]
    todo += mancanti[: args.max_details]
    log(f"portale: {len(todo)} dettagli da scaricare ({len(mancanti)} chiusi senza dettaglio in cache)")

    lock, t0 = threading.Lock(), time.time()

    def work(it):
        if time.time() - t0 > args.time_budget:
            return
        try:
            time.sleep(0.15)
            d = portale_dettaglio(it["id"])
        except Exception:
            return
        with lock:
            cache[str(it["id"])] = d

    with ThreadPoolExecutor(4) as ex:
        list(ex.map(work, todo))

    rows = []
    for it in items:
        d = cache.get(str(it["id"]))
        if d and d.get("err"):
            d = None
        tipo = "Bando" if it["tipo"] == 1 else "Avviso"
        r = riga_vuota()
        cat = (d or {}).get("categoria", "")
        v = importo((d or {}).get("importo", "")) if d else None
        og = (d or {}).get("oggetto", "")
        if len(og) > 650:
            og = og[:650].rsplit(" ", 1)[0] + "…"
        r.update(id=str(it["id"]), tipo=tipo, stato_portale=it["stato"], titolo=fix(it["titolo"]),
                 struttura=fix(it["struttura"]), categoria="" if cat == "- Non Selezionata -" else cat,
                 gruppo=classifica(tipo, cat, it["titolo"], bool(d)), scadenza=txt_date(it["scadenza"]),
                 scad_iso=iso(it["scadenza"]), pubblicazione=(d or {}).get("pubblicazione", ""),
                 pub_iso=iso((d or {}).get("pubblicazione", "")), importo=v, importo_txt=importo_txt(v) if d else "",
                 destinatari=(d or {}).get("destinatari", ""),
                 materie=" | ".join(fix(m) for m in (d or {}).get("materie", [])), oggetto=og,
                 allegati=(d or {}).get("nattach", 0), dettaglio=bool(d),
                 link=f"{PORTAL}/Public/Dettaglio?idAtto={it['id']}", fonte="Portale Bandi Regione Veneto")
        rows.append(r)
    return rows


# ------------------------------------------------------------------ fonte 2: calendario FESR
def fonte_fesr(rows_portale):
    page = get(COESIONE + "/fesr/calendario-inviti").text
    d = BeautifulSoup(page, "html.parser")
    tab = d.find("table")
    voci = []
    if tab:
        for tr in tab.find_all("tr"):
            c = [re.sub(r"\s+", " ", td.get_text(" ", strip=True)) for td in tr.find_all(["td", "th"])]
            if len(c) >= 6 and re.search(r"\d\.\d", c[1]):
                voci.append(dict(oggetto=c[0], codice=c[1], direzione=c[2], mese=c[3], modalita=c[4], benef=c[5]))
    # CSV con importi e obiettivi
    importi = {}
    link = next((urljoin(COESIONE, a["href"]) for a in d.find_all("a", href=True)
                 if "documents" in a["href"] and re.search(r"cronoprogramma.*csv|csv", a.get_text(" "), re.I)), None)
    if link:
        raw = get(link).content
        txt = raw.decode("utf-8-sig", "replace")
        for r in csv.reader(io.StringIO(txt)):
            if len(r) >= 8 and re.match(r"\s*\d\.\d", r[0]):
                code = re.match(r"\s*(\d\.\d+(?:\.\d+)?)", r[0]).group(1)
                importi.setdefault(code, []).append((importo(r[6]), re.sub(r"\s+", " ", r[4]), r[0].strip()))
    out = []
    this_month = TODAY.replace(day=1).isoformat()
    for v in voci:
        code = re.match(r"\s*(\d\.\d+(?:\.\d+)?)", v["codice"]).group(1)
        ini, fine = parse_finestra(v["mese"])
        if not ini or fine < this_month:
            continue
        # già pubblicato nel portale? (stesso codice azione, e stesso sub-bando se indicato)
        sub = re.search(r"sub\.?\s*([a-z])\b|\b([a-z])\)", v["codice"], re.I)
        sub = (sub.group(1) or sub.group(2)).lower() if sub else None

        def gia_pubblicato(p):
            if not re.search(r"(?<![\d.])" + re.escape(code) + r"(?![\d])", p["titolo"]):
                return False
            if not (p["pub_iso"] >= ini[:4] + "-01-01" or p["stato_portale"] == "PUBBLICATO"):
                return False
            ts = re.findall(r"sub\.?\s*([a-z])\b", p["titolo"], re.I)
            return (not sub) or (not ts) or sub in [x.lower() for x in ts]

        if any(gia_pubblicato(p) for p in rows_portale):
            continue
        imp, descr, arg = (importi.get(code) or [(None, "", "")]).pop(0) if importi.get(code) else (None, "", "")
        r = riga_vuota()
        r.update(id=f"fesr-{code}-{ini}", tipo="Previsto", stato_portale="PREVISTO",
                 titolo=f"PR Veneto FESR 2021-27 · Azione {v['codice']} · {v['oggetto']}", struttura=v["direzione"],
                 categoria=v["modalita"], gruppo="Finanziamenti", pubblicazione=v["mese"], pub_iso=ini,
                 apertura_prevista=v["mese"], importo=imp, importo_txt=importo_txt(imp) if imp else "n.d.",
                 destinatari="Impresa" if re.search(r"imprese|PMI", v["benef"], re.I) else "Ente",
                 materie="FESR", oggetto=(arg + (" — " + descr if descr else "") + " Beneficiari: " + v["benef"])[:650],
                 dettaglio=True, link=COESIONE + "/fesr/calendario-inviti",
                 fonte="Calendario inviti FESR 2021-27")
        out.append(r)
    return out


# ------------------------------------------------------------------ fonte 3: calendari FSE+
def decodifica_csv(raw):
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("latin-1")


def fonte_fse(rows_portale, ultimi=2):
    page = get(COESIONE + "/fse/calendario-inviti").text
    d = BeautifulSoup(page, "html.parser")
    links = [(re.sub(r"\s+", " ", a.get_text(" ", strip=True)), urljoin(COESIONE, a["href"]))
             for a in d.find_all("a", href=True) if "documents" in a["href"] and re.search(r"csv", a.get_text(" "), re.I)]
    out, this_month = [], TODAY.replace(day=1).isoformat()
    tit_portale = [(p, toks(p["titolo"])) for p in rows_portale if p["stato_portale"] == "PUBBLICATO" or p["pub_iso"] >= f"{TODAY.year}-01-01"]
    for nome, url in links[:ultimi]:
        rows = list(csv.reader(io.StringIO(decodifica_csv(get(url).content)), delimiter=";"))
        hi = next((i for i, r in enumerate(rows) if "PROGRAMMA" in [x.strip().upper() for x in r]), None)
        if hi is None:
            continue
        h = [x.strip().lower() for x in rows[hi]]

        def col(*keys):
            return next((i for i, x in enumerate(h) if all(k in x for k in keys)), None)

        c_arg, c_ric, c_imp, c_dat = col("argomento"), col("richiedenti"), col("importo"), col("data di apertura")
        c_sra, c_desc = col("sra"), (col("obiettivo specifico") or 0) + 1
        for r in rows[hi + 1:]:
            if len(r) <= max(c_arg, c_dat) or not r[0].strip().startswith("2021IT"):
                continue
            ini, fine = parse_finestra(r[c_dat])
            if not ini or fine < this_month:
                continue
            arg = fix(r[c_arg]).rstrip("* ").strip()
            if any(jaccard(toks(arg), pt) >= 0.35 for _, pt in tit_portale):
                continue
            imp = importo(r[c_imp])
            rr = riga_vuota()
            rr.update(id=f"fse-{re.sub(r'[^a-z0-9]+', '-', arg.lower())[:40]}-{ini}", tipo="Previsto",
                      stato_portale="PREVISTO", titolo=f"PR Veneto FSE+ 2021-27 · {arg}",
                      struttura="Regione del Veneto · Autorità di Gestione FSE+", categoria="Invito FSE+",
                      gruppo="Finanziamenti", pubblicazione=fix(r[c_dat]), pub_iso=ini,
                      apertura_prevista=fix(r[c_dat]), importo=imp, importo_txt=importo_txt(imp),
                      destinatari="Ente", materie=fix(r[c_sra]) if c_sra is not None and len(r) > c_sra and r[c_sra].strip() else "Formazione",
                      oggetto=(fix(r[c_desc]) + " Richiedenti: " + fix(r[c_ric]))[:650], dettaglio=True,
                      link=COESIONE + "/fse/calendario-inviti", fonte="Calendario inviti FSE+")
            out.append(rr)
    return out


# ------------------------------------------------------------------ fonte 4: foglio Opportunità (Coesione)
def fonte_coesione(rows_portale):
    rows = list(csv.reader(io.StringIO(get(SHEET).text)))
    hdr = rows[0]
    out = []
    tit = [(p, toks(p["titolo"])) for p in rows_portale]
    for r in rows[1:]:
        c = dict(zip(hdr, r))
        if c.get("Categoria") != "bandi":
            continue
        scad = iso(c.get("Scadenza", ""))
        if not scad or scad < TODAY.isoformat():
            continue
        az = re.findall(r"\d\.\d+(?:\.\d+)?", c.get("Azione", ""))
        t_c = toks(c["TitoloAlto"] + " " + c.get("DescrizioneBando", "")[:200])
        if any((az and az[0] in p["titolo"]) or jaccard(t_c, pt) >= 0.3 for p, pt in tit):
            continue
        link = next((c[k] for k in ("SchedaLink", "Link2", "Form link") if c.get(k, "").startswith("http")),
                    COESIONE + "/opportunit%C3%A0")
        rr = riga_vuota()
        rr.update(id=f"coes-{abs(hash(c['TitoloAlto'] + c.get('Azione', '') + scad)) % 10**8}", tipo="Bando",
                  stato_portale="PUBBLICATO", titolo=fix(f"{c['TitoloAlto']} · {c.get('Azione', '')} · {c.get('DescrizioneBando', '')[:80]}").strip(" ·"),
                  struttura="Regione del Veneto · Programmazione FESR/FSE+", categoria=c.get("Asse", "")[:60],
                  gruppo="Finanziamenti", scadenza=txt_date(c["Scadenza"]), scad_iso=scad,
                  pubblicazione=txt_date(c.get("Pubblicazione", "")), pub_iso=iso(c.get("Pubblicazione", "")),
                  materie="FESR", oggetto=fix(c.get("DescrizioneBando", ""))[:650], dettaglio=True, importo_txt="n.d.",
                  link=link, fonte="Veneto Coesione (Opportunità)")
        out.append(rr)
    return out



# ------------------------------------------------------------------ export Excel
def stato_calcolato(r):
    if r["stato_portale"] == "PREVISTO":
        return "Previsto"
    if r["stato_portale"] == "PUBBLICATO" and (not r["scad_iso"] or r["scad_iso"] >= TODAY.isoformat()):
        if r["scad_iso"] and r["scad_iso"] <= (TODAY + dt.timedelta(days=14)).isoformat():
            return "In scadenza"
        if r["pub_iso"] and r["pub_iso"] >= (TODAY - dt.timedelta(days=14)).isoformat():
            return "Nuovo"
        return "Aperto"
    return "Chiuso"


def scrivi_xlsx(pack, path):
    """Dataset completo in Excel (foglio Bandi con filtri + foglio Info)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        log("openpyxl non installato: salto l'export Excel (pip install openpyxl)")
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Bandi"
    head = ["Stato", "Titolo", "Atto", "Tipologia", "Struttura", "Categoria", "Dotazione (€)", "Pubblicazione",
            "Scadenza", "Apertura prevista", "Destinatari", "Materie", "Fonte", "Oggetto", "Link"]
    ws.append(head)
    for r in pack["rows"]:
        def d(s):
            return dt.date.fromisoformat(s) if s else None
        prev = r["stato_portale"] == "PREVISTO"
        ws.append([stato_calcolato(r), r["titolo"], r["tipo"], r["gruppo"], r["struttura"], r["categoria"],
                   round(r["importo"]) if r["importo"] else None, None if prev else d(r["pub_iso"]), d(r["scad_iso"]),
                   r["apertura_prevista"] or None, r["destinatari"], r["materie"], r["fonte"], r["oggetto"], None])
        c = ws.cell(row=ws.max_row, column=15)
        if r["link"]:
            c.value, c.hyperlink = "Apri", r["link"]
            c.font = Font(color="0563C1", underline="single")
        ws.cell(row=ws.max_row, column=7).number_format = "#,##0"
        for col in (8, 9):
            ws.cell(row=ws.max_row, column=col).number_format = "DD/MM/YYYY"
    fill = PatternFill("solid", fgColor="E8E6DC")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), fill
    for i, w in enumerate([13, 70, 9, 18, 38, 22, 14, 14, 14, 20, 12, 30, 28, 80, 9], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet("Info")
    for row in (["Campo", "Valore"], ["Dati generati il", pack["generato"]], ["Righe", len(pack["rows"])],
                ["Stato calcolato al", TODAY.isoformat()],
                ["Fonti", "; ".join(f["fonte"] + ("" if f["ok"] else " (non aggiornata)") for f in pack["fonti"])],
                ["Nota", "Il testo pubblicato dal portale non ha valore ufficiale: fa fede il Bollettino Ufficiale della Regione del Veneto."]):
        info.append(row)
    for c in info[1]:
        c.font, c.fill = Font(bold=True), fill
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 20, 100
    wb.save(path)
    return True

# ------------------------------------------------------------------ build
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=HERE)
    ap.add_argument("--max-details", type=int, default=150, help="dettagli di bandi chiusi da scaricare per esecuzione")
    ap.add_argument("--time-budget", type=int, default=900, help="secondi massimi per lo scaricamento dei dettagli")
    ap.add_argument("--no-refresh-open", action="store_true", help="non riscaricare i dettagli degli atti aperti")
    args = ap.parse_args()

    cdir = os.path.join(args.out, "cache"); ddir = os.path.join(args.out, "data")
    os.makedirs(cdir, exist_ok=True); os.makedirs(ddir, exist_ok=True)
    cpath = os.path.join(cdir, "detail_cache.json")
    cache = json.load(open(cpath, encoding="utf-8")) if os.path.exists(cpath) else {}

    rows, stato_fonti = [], []
    try:
        rp = fonte_portale(args, cache)
        rows += rp
        stato_fonti.append({"fonte": "Portale Bandi Regione Veneto", "ok": True, "righe": len(rp)})
    except Exception as e:
        log("ERRORE portale:", e); rp = []
        stato_fonti.append({"fonte": "Portale Bandi Regione Veneto", "ok": False, "errore": str(e)[:120]})
    json.dump(cache, open(cpath, "w", encoding="utf-8"), ensure_ascii=False)

    for nome, fn in (("Calendario inviti FESR 2021-27", fonte_fesr), ("Calendario inviti FSE+", fonte_fse),
                     ("Veneto Coesione (Opportunità)", fonte_coesione)):
        try:
            extra = fn(rp)
            rows += extra
            stato_fonti.append({"fonte": nome, "ok": True, "righe": len(extra)})
            log(f"{nome}: {len(extra)} righe")
        except Exception as e:
            log(f"ERRORE {nome}:", e)
            stato_fonti.append({"fonte": nome, "ok": False, "errore": str(e)[:120]})

    # se il portale non è raggiungibile, riusa l'ultimo dataset invece di pubblicare un dataset monco
    prev = os.path.join(ddir, "bandi_veneto.json")
    if not rp and os.path.exists(prev):
        log("portale non raggiungibile: mantengo l'ultimo dataset salvato")
        sys.exit(1)

    pack = {"generato": dt.datetime.now().strftime("%Y-%m-%d %H:%M"), "fonti": stato_fonti, "rows": rows}
    json.dump(pack, open(prev, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))

    if scrivi_xlsx(pack, os.path.join(ddir, "bandi_veneto.xlsx")):
        log("Excel: data/bandi_veneto.xlsx")
    # copia dei dati accanto a docs/index.html: l'app online li rilegge a ogni apertura
    import shutil
    os.makedirs(os.path.join(args.out, "docs", "data"), exist_ok=True)
    for nome in ("bandi_veneto.json", "bandi_veneto.xlsx"):
        if os.path.exists(os.path.join(ddir, nome)):
            shutil.copyfile(os.path.join(ddir, nome), os.path.join(args.out, "docs", "data", nome))
    tpl = open(os.path.join(HERE, "standalone_template.html"), encoding="utf-8").read()
    blob = json.dumps(pack, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    page = tpl.replace("__BV_DATA__", blob)
    open(os.path.join(args.out, "bandi_veneto.html"), "w", encoding="utf-8").write(page)
    os.makedirs(os.path.join(args.out, "docs"), exist_ok=True)       # per GitHub Pages
    open(os.path.join(args.out, "docs", "index.html"), "w", encoding="utf-8").write(page)
    n_aperti = sum(1 for r in rows if r["stato_portale"] == "PUBBLICATO")
    log(f"fatto: {len(rows)} righe ({n_aperti} aperte, {sum(1 for r in rows if r['stato_portale'] == 'PREVISTO')} previste)")


if __name__ == "__main__":
    main()
